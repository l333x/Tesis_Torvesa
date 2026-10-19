"""Análisis provisional de Ingresos_Torvesa_2025.xlsx. Uso: python -I analisis.py <xlsx> <out_dir>
Guarda resultados.json y figuras PNG en out_dir."""
import json, sys, statistics as st
from datetime import date, datetime
from pathlib import Path
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

xlsx, out = Path(sys.argv[1]), Path(sys.argv[2]); out.mkdir(parents=True, exist_ok=True)
wb = openpyxl.load_workbook(xlsx, data_only=True)
wbf = openpyxl.load_workbook(xlsx)  # fórmulas
MES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

# ---- proyectos
wp = wb["Proyectos"]; hdr = [c.value for c in wp[4]]
proy = {}
for r in wp.iter_rows(min_row=5, max_row=20, values_only=True):
    d = dict(zip(hdr, r)); proy[d["Proyecto"]] = d
ip = wb["Ingresos por proyecto"]; ipf = wbf["Ingresos por proyecto"]
for r in range(5, 21):
    pid = ip.cell(r, 1).value
    proy[pid]["mes"] = [ip.cell(r, c).value or 0 for c in range(5, 17)]
    proy[pid]["formulas_en_meses"] = sum(1 for c in range(5, 17) if isinstance(ipf.cell(r, c).value, str) and ipf.cell(r, c).value.startswith("="))
wc = wb["Ingresos por cliente"]
cli = {}
for r in range(5, 17):
    v = [wc.cell(r, c).value for c in range(1, 6)]
    cli[v[0]] = dict(nombre=v[1], segmento=v[2], origen=v[3], zona=v[4], mes=[wc.cell(r, c).value or 0 for c in range(6, 18)])
wr = wb["Resumen mensual"]
res = [[wr.cell(r, c).value for c in range(1, 15)] for r in range(5, 17)]

R = {}
tot = [sum(p["mes"][m] for p in proy.values()) for m in range(12)]
pub = [sum(p["mes"][m] for p in proy.values() if p["Origen de la demanda"].startswith("Obra")) for m in range(12)]
T = sum(tot)
R["total"] = T; R["mensual"] = tot; R["publica"] = pub
R["promedio"] = T / 12; R["cv"] = st.pstdev(tot) / (T / 12)
R["trimestres"] = [sum(tot[i:i + 3]) for i in (0, 3, 6, 9)]
R["s1"], R["s2"] = sum(tot[:6]), sum(tot[6:])

# ---- controles de consistencia
ctrl = []
def chk(nombre, ok, det=""): ctrl.append(dict(control=nombre, ok=bool(ok), detalle=det))
chk("Suma por proyectos = suma por clientes", abs(T - sum(sum(c["mes"]) for c in cli.values())) < 0.05)
chk("Resumen mensual = suma de proyectos", all(abs(res[m][1] - tot[m]) < 0.05 for m in range(12)))
bad = []
for k, p in proy.items():
    M = p["Monto sin IVA"]; e24 = p["Ejecutado hasta 31-12-2024"] or 0; i25 = p["Ingreso reconocido 2025"]; s26 = p["Saldo por ejecutar en 2026"] or 0
    if abs(M - e24 - i25 - s26) > 0.05: bad.append(f"{k}: monto ≠ 2024+2025+saldo")
    if abs(i25 - sum(p["mes"])) > 0.05: bad.append(f"{k}: ingreso 2025 ≠ suma mensual")
    if abs((p["IVA 15 %"] or 0) - 0.15 * M) > 0.05: bad.append(f"{k}: IVA ≠ 15 %")
    if abs((p["Anticipo (USD)"] or 0) - (p["Anticipo %"] or 0) * M) > 0.05: bad.append(f"{k}: anticipo ≠ % × monto")
    ini, fin = p["Inicio"].date(), p["Fin real (con prórroga)"].date()
    for m in range(12):
        if p["mes"][m] > 0 and not (date(2025, m + 1, 28) >= ini.replace(day=1) and date(2025, m + 1, 1) <= fin):
            bad.append(f"{k}: ingreso en {MES[m]} fuera de su plazo")
    plazo = (p["Fin contractual"] - p["Inicio"]).days
    if abs(plazo - p["Plazo contractual (días)"]) > 1: bad.append(f"{k}: plazo {p['Plazo contractual (días)']} vs fechas {plazo}")
chk("Montos, IVA, anticipos, fechas y plazos de cada contrato", not bad, "; ".join(bad))
for k in ("P08", "P12"):
    p = proy[k]; chk(f"{k}: subcontrato ≤ 30 % del contrato principal", p["% del principal"] <= 0.30, f"{p['% del principal']:.1%}")
R["controles"] = ctrl
R["celdas_mensuales_tecleadas"] = sum(16 * 12 - 0 for _ in [0]) - sum(p["formulas_en_meses"] for p in proy.values())

# ---- clientes, segmentos
ct = {k: sum(c["mes"]) for k, c in cli.items()}
orden = sorted(ct, key=ct.get, reverse=True)
R["clientes"] = [dict(cod=k, nombre=cli[k]["nombre"], segmento=cli[k]["segmento"], total=ct[k], part=ct[k] / T,
                      proyectos=sum(1 for p in proy.values() if p["Código cliente"] == k)) for k in orden]
R["hhi"] = sum((v / T * 100) ** 2 for v in ct.values())
rec = [k for k in ct if sum(1 for p in proy.values() if p["Código cliente"] == k) > 1]
R["recurrentes"] = dict(clientes=rec, n=len(rec), ingreso=sum(ct[k] for k in rec), part=sum(ct[k] for k in rec) / T)
seg = {}
for k, c in cli.items(): seg[c["segmento"]] = seg.get(c["segmento"], 0) + ct[k]
R["segmentos"] = sorted(seg.items(), key=lambda x: -x[1])

# ---- contratos (lo comercial)
ps = sorted(proy.values(), key=lambda p: p["Inicio"])
montos = [p["Monto sin IVA"] for p in ps]
R["contratos"] = dict(n=len(ps), n_firmados_2025=sum(1 for p in ps if p["Inicio"].year == 2025),
                      monto_total=sum(montos), promedio=st.mean(montos), mediana=st.median(montos),
                      minimo=min(montos), maximo=max(montos),
                      suma_alzada=sum(1 for p in ps if "suma alzada" in p["Modalidad"]),
                      plazo_mediano=st.median(p["Plazo contractual (días)"] for p in ps),
                      prorroga_total=sum(p["Prórroga por paralizaciones (días)"] or 0 for p in ps),
                      con_prorroga=sum(1 for p in ps if p["Prórroga por paralizaciones (días)"]),
                      inicios_por_mes=[sum(1 for p in ps if p["Inicio"].year == 2025 and p["Inicio"].month == m + 1) for m in range(12)],
                      saldo_2026=sum(p["Saldo por ejecutar en 2026"] or 0 for p in ps),
                      anticipo_prom=st.mean(p["Anticipo %"] for p in ps))
R["proyectos"] = [dict(id=p["Proyecto"], cliente=p["Código cliente"], origen=p["Origen de la demanda"], modalidad=p["Modalidad"],
                       monto=p["Monto sin IVA"], ing2025=p["Ingreso reconocido 2025"], inicio=p["Inicio"].date().isoformat(),
                       fin=p["Fin real (con prórroga)"].date().isoformat(), plazo=p["Plazo contractual (días)"],
                       anticipo=p["Anticipo %"], pago=p["Forma de pago"], prorroga=p["Prórroga por paralizaciones (días)"])
                  for p in ps]

# ---- señales de que la serie mensual es modelada
obs = [x[7] for x in res]; usado = [x[8] for x in res]; sector = [x[9] for x in res]
def corr(a, b):
    ma, mb = st.mean(a), st.mean(b); num = sum((x - ma) * (y - mb) for x, y in zip(a, b))
    return num / (sum((x - ma) ** 2 for x in a) * sum((y - mb) ** 2 for y in b)) ** 0.5
R["senales"] = dict(corr_obs_usado=corr(obs, usado), corr_obs_sector=corr(obs, sector),
                    sri_dic_original=res[11][12], sri_dic_usado=res[11][13],
                    trabajadores=wr["B27"].value, ref_mediana=wr["B29"].value)
R["ingreso_trab"] = T / wr["B27"].value

json.dump(R, open(out / "resultados.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)

# ---- figuras
AZ, CO = "#0B1F3A", "#B87333"
fig, ax = plt.subplots(figsize=(9, 3.6))
priv = [t - p for t, p in zip(tot, pub)]
ax.bar(MES, [x / 1000 for x in priv], color=AZ, label="Privada")
ax.bar(MES, [x / 1000 for x in pub], bottom=[x / 1000 for x in priv], color=CO, label="Obra pública (subcontrato)")
ax.axhline(T / 12 / 1000, ls="--", color="grey", lw=1); ax.text(11.4, T / 12 / 1000 + 3, "promedio", ha="right", fontsize=8, color="grey")
ax.set_ylabel("miles de USD (sin IVA)"); ax.legend(frameon=False, fontsize=8); ax.spines[["top", "right"]].set_visible(False)
fig.tight_layout(); fig.savefig(out / "fig1_mensual.png", dpi=180); plt.close(fig)

fig, ax = plt.subplots(figsize=(9, 3.6))
vals = [ct[k] / 1000 for k in orden]; acum = [sum(vals[:i + 1]) / (T / 1000) * 100 for i in range(len(vals))]
ax.bar(orden, vals, color=[CO if k in rec else AZ for k in orden]); ax.set_ylabel("miles de USD")
ax2 = ax.twinx(); ax2.plot(orden, acum, color="grey", marker="o", ms=3); ax2.set_ylim(0, 105); ax2.set_ylabel("% acumulado")
ax.spines[["top"]].set_visible(False); ax.set_title("Ingreso 2025 por cliente (cobre = cliente con más de un contrato)", fontsize=9)
fig.tight_layout(); fig.savefig(out / "fig2_clientes.png", dpi=180); plt.close(fig)

fig, ax = plt.subplots(figsize=(9, 4.2))
d0 = date(2024, 10, 1)
for i, p in enumerate(ps):
    a, b = (p["Inicio"].date() - d0).days, (p["Fin real (con prórroga)"].date() - d0).days
    ax.barh(i, b - a, left=a, color=CO if p["Origen de la demanda"].startswith("Obra") else AZ)
    ax.text(b + 3, i, f'{p["Proyecto"]} · {p["Monto sin IVA"]/1000:,.0f} mil', va="center", fontsize=7)
ax.set_yticks([]); ax.invert_yaxis()
tk = [date(2024, 10, 1), date(2025, 1, 1), date(2025, 4, 1), date(2025, 7, 1), date(2025, 10, 1), date(2026, 1, 1)]
ax.set_xticks([(t - d0).days for t in tk]); ax.set_xticklabels([t.strftime("%b-%y") for t in tk], fontsize=8)
ax.spines[["top", "right", "left"]].set_visible(False); ax.set_title("Contratos en el tiempo (cobre = subcontrato de obra pública)", fontsize=9)
fig.tight_layout(); fig.savefig(out / "fig3_contratos.png", dpi=180); plt.close(fig)
print(json.dumps({k: R[k] for k in ("total", "promedio", "cv", "trimestres", "s1", "s2", "hhi", "recurrentes", "segmentos", "contratos", "senales", "ingreso_trab", "controles")}, ensure_ascii=False, indent=1, default=str))
