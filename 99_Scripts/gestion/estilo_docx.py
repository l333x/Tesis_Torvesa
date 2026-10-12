"""Estilo común de los documentos Word de la tesis (azul marino + cobre, Poppins/Lato) y exportación a PDF.

Uso:
    from estilo_docx import Doc
    d = Doc("Título", "Subtítulo", "2026-10-11")
    d.h1("Sección"); d.p("Texto con **negrita** y `código`."); d.tabla([["A","B"],["1","2"]])
    d.comentario(parrafo, "Víctor: decide X", autor="Claude")
    d.guardar(ruta_docx, pdf=True)
También convierte Markdown simple: d.markdown(texto_md)
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x0B, 0x1F, 0x3A)
COBRE = RGBColor(0xB8, 0x73, 0x33)
GRIS = RGBColor(0x55, 0x5E, 0x6B)
HEX_AZUL, HEX_COBRE, HEX_CLARO, HEX_COBRE_CLARO = "0B1F3A", "B87333", "EEF2F7", "F7EDE3"
TITULOS, CUERPO = "Poppins", "Lato"
# Los emoji obligan a Word a incrustar Segoe UI Emoji (~4 MB por PDF): se reemplazan por texto o símbolos simples
EMOJI = {"⚠️": "", "⚠": "", "✂️": "(recortable)", "✅": "✓", "⛔": "✗", "❌": "✗", "⏳": "(pendiente)",
         "🔎": "(por verificar)", "🏠": "(dato propio)", "⏸️": "(en pausa)", "ℹ️": "", "🚨": "", "✔": "✓",
         "️": ""}


def _sombra(celda, hex_color: str) -> None:
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def _fuente(run, nombre: str) -> None:
    run.font.name = nombre
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), nombre)


def _borde_inferior(parrafo, hex_color: str, grosor: int = 12) -> None:
    ppr = parrafo._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(grosor)); b.set(qn("w:space"), "4"); b.set(qn("w:color"), hex_color)
    bdr.append(b); ppr.append(bdr)


class Doc:
    def __init__(self, titulo: str, subtitulo: str = "", fecha: str = "", horizontal: bool = False):
        self.d = Document()
        sec = self.d.sections[0]
        if horizontal:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = sec.page_height, sec.page_width
        for m in ("left_margin", "right_margin"):
            setattr(sec, m, Cm(2.2))
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        self._estilos()
        self._pie(titulo)
        self.portada(titulo, subtitulo, fecha)

    # ---------- estilos base ----------
    def _estilos(self) -> None:
        st = self.d.styles
        n = st["Normal"]; n.font.name = CUERPO; n.font.size = Pt(10.5); n.font.color.rgb = RGBColor(0x22, 0x2A, 0x35)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), CUERPO)
        n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.15
        for nivel, tam, color in ((1, 17, AZUL), (2, 13, COBRE), (3, 11.5, AZUL)):
            h = st[f"Heading {nivel}"]
            h.font.name = TITULOS; h.font.size = Pt(tam); h.font.bold = True; h.font.color.rgb = color
            rf = h.element.rPr.rFonts
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):  # el tema pisaría a Poppins
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), TITULOS)
            h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10); h.paragraph_format.space_after = Pt(4)
        for lista in ("List Bullet", "List Number"):
            st[lista].font.name = CUERPO; st[lista].font.size = Pt(10.5)

    def _pie(self, titulo: str) -> None:
        pie = self.d.sections[0].footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = pie.add_run(f"{titulo}  ·  Tesis Torvesa · UPS  ·  pág. ")
        r.font.size = Pt(8); r.font.color.rgb = GRIS; _fuente(r, CUERPO)
        r2 = pie.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r2._r.append(f)
            else:
                i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = txt; r2._r.append(i)
        r2.font.size = Pt(8); r2.font.color.rgb = GRIS

    # ---------- bloques ----------
    def portada(self, titulo: str, subtitulo: str, fecha: str) -> None:
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0); _sombra(c, HEX_AZUL)
        p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(10)
        r = p.add_run("TESIS TORVESA · UPS"); r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = COBRE; _fuente(r, TITULOS)
        p2 = c.add_paragraph(); r = p2.add_run(titulo); r.font.size = Pt(22); r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); _fuente(r, TITULOS)
        if subtitulo:
            p3 = c.add_paragraph(); r = p3.add_run(subtitulo); r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0xE6, 0xC3, 0xA0); _fuente(r, CUERPO)
        p4 = c.add_paragraph(); p4.paragraph_format.space_after = Pt(10)
        r = p4.add_run(fecha or ""); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xCF, 0xD8, 0xE3); _fuente(r, CUERPO)
        self.d.add_paragraph()

    def h1(self, txt: str):
        p = self.d.add_heading(txt, level=1); _borde_inferior(p, HEX_COBRE, 10); return p

    def h2(self, txt: str):
        return self.d.add_heading(txt, level=2)

    def h3(self, txt: str):
        return self.d.add_heading(txt, level=3)

    def _hipervinculo(self, p, texto: str, url: str) -> None:
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
        h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
        r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color"); c.set(qn("w:val"), HEX_COBRE); rpr.append(c)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        r.append(rpr); t = OxmlElement("w:t"); t.text = texto; t.set(qn("xml:space"), "preserve"); r.append(t)
        h.append(r); p._p.append(h)

    def _inline(self, p, texto: str, base_bold: bool = False) -> None:
        # **negrita**, *cursiva*, `código`, [texto](url): http -> hipervínculo; ruta relativa -> solo texto
        for e, s in EMOJI.items():
            texto = texto.replace(e, s)
        for trozo in re.split(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|<https?://[^>]+>)", texto):
            if not trozo:
                continue
            m = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", trozo) or re.match(r"^<((https?://[^>]+))>$", trozo)
            if m:
                txt, url = m.group(1), m.group(2)
                if url.startswith("http"):
                    self._hipervinculo(p, txt.strip("`"), url)
                else:
                    self._inline(p, txt)
                continue
            if trozo.startswith("**"):
                r = p.add_run(trozo[2:-2]); r.bold = True
            elif trozo.startswith("`"):
                r = p.add_run(trozo[1:-1]); _fuente(r, "Consolas"); r.font.size = Pt(9.5); r.font.color.rgb = AZUL
            elif trozo.startswith("*") and len(trozo) > 2:
                r = p.add_run(trozo[1:-1]); r.italic = True
            else:
                r = p.add_run(trozo); r.bold = base_bold or None

    def p(self, texto: str = "", estilo: str | None = None, negrita: bool = False):
        par = self.d.add_paragraph(style=estilo)
        self._inline(par, texto, negrita)
        return par

    def vineta(self, texto: str, numerada: bool = False, nivel: int = 0):
        par = self.p(texto, "List Number" if numerada else "List Bullet")
        if nivel:
            par.paragraph_format.left_indent = Cm(0.63 * (nivel + 1))
        return par

    def check(self, texto: str, hecho: bool):
        par = self.d.add_paragraph()
        r = par.add_run("☑ " if hecho else "☐ "); r.font.color.rgb = COBRE if hecho else GRIS; r.bold = True
        self._inline(par, texto)
        return par

    def aviso(self, texto: str, tipo: str = "info"):
        """Recuadro destacado (info | alerta | ok)."""
        colores = {"info": (HEX_CLARO, AZUL), "alerta": (HEX_COBRE_CLARO, COBRE), "ok": ("E8F3EC", RGBColor(0x2E, 0x6B, 0x45))}
        fondo, borde = colores.get(tipo, colores["info"])
        t = self.d.add_table(rows=1, cols=1); c = t.cell(0, 0); _sombra(c, fondo)
        lineas = texto.split("\n")
        self._inline(c.paragraphs[0], lineas[0])
        for l in lineas[1:]:
            self._inline(c.add_paragraph(), l)
        self.d.add_paragraph()
        return c.paragraphs[0]

    def tabla(self, filas: list[list[str]], anchos: list[float] | None = None, encabezado: bool = True):
        t = self.d.add_table(rows=len(filas), cols=len(filas[0]))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if anchos:  # sin diseño fijo, Word reparte el ancho a su gusto e ignora los anchos
            t.autofit = False
            lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
            for j, w in enumerate(anchos):
                t.columns[j].width = Cm(w)
        for i, fila in enumerate(filas):
            for j, val in enumerate(fila):
                c = t.cell(i, j); c.paragraphs[0].text = ""
                par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(1)
                if i == 0 and encabezado:
                    _sombra(c, HEX_AZUL)
                    r = par.add_run(str(val)); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9.5); _fuente(r, TITULOS)
                else:
                    if i % 2 == 0:
                        _sombra(c, HEX_CLARO)
                    self._inline(par, str(val))
                    for r in par.runs:
                        r.font.size = Pt(9.5)
                if anchos:
                    c.width = Cm(anchos[j])
        # bordes finos cobre
        tblpr = t._tbl.tblPr; bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{lado}"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "C9A27E" if lado.startswith("inside") else HEX_COBRE); bordes.append(b)
        tblpr.append(bordes)
        self.d.add_paragraph()
        return t

    def comentario(self, parrafo, texto: str, autor: str = "Claude", iniciales: str = "CL"):
        """Comentario real de Word anclado al párrafo (python-docx >= 1.2)."""
        runs = parrafo.runs or [parrafo.add_run("")]
        return self.d.add_comment(runs, text=texto, author=autor, initials=iniciales)

    def salto(self):
        self.d.add_page_break()

    def imagen(self, ruta, ancho_cm: float = 16.0, pie: str = ""):
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ruta), width=Cm(ancho_cm))
        if pie:
            c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(pie); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = GRIS
        return p

    # ---------- Markdown simple -> docx ----------
    def markdown(self, md: str) -> None:
        md = re.sub(r"^---\n.*?\n---\n", "", md, count=1, flags=re.S)  # frontmatter
        lineas = md.split("\n"); i = 0
        while i < len(lineas):
            l = lineas[i]
            if l.startswith("```"):
                bloque = []; i += 1
                while i < len(lineas) and not lineas[i].startswith("```"):
                    bloque.append(lineas[i]); i += 1
                t = self.d.add_table(rows=1, cols=1); c = t.cell(0, 0); _sombra(c, "F3F5F8")
                c.paragraphs[0].text = ""
                for k, bl in enumerate(bloque):
                    par = c.paragraphs[0] if k == 0 else c.add_paragraph()
                    par.paragraph_format.space_after = Pt(0)
                    r = par.add_run(bl); _fuente(r, "Consolas"); r.font.size = Pt(9)
                self.d.add_paragraph(); i += 1; continue
            if l.startswith("|") and i + 1 < len(lineas) and re.match(r"^\|[\s:|-]+\|$", lineas[i + 1].strip()):
                filas = []
                while i < len(lineas) and lineas[i].startswith("|"):
                    if not re.match(r"^\|[\s:|-]+\|$", lineas[i].strip()):
                        filas.append([x.strip() for x in lineas[i].strip().strip("|").split("|")])
                    i += 1
                n = max(len(f) for f in filas); filas = [f + [""] * (n - len(f)) for f in filas]
                self.tabla(filas); continue
            if l.startswith("> "):
                bloque = []
                while i < len(lineas) and lineas[i].startswith(">"):
                    bloque.append(re.sub(r"^>\s?", "", lineas[i])); i += 1
                bloque = [re.sub(r"^\[![\w-]+\][-+]?\s*", "", b) for b in bloque]
                self.aviso("\n".join(b for b in bloque if b.strip()), "alerta" if "⚠" in "".join(bloque) or "[!warning" in l or "[!danger" in l else "info")
                continue
            m = re.match(r"^(#{1,4})\s+(.*)", l)
            if m:
                {1: self.h1, 2: self.h1, 3: self.h2, 4: self.h3}[len(m.group(1))](m.group(2))
            elif re.match(r"^\s*- \[[ xX]\] ", l):
                self.check(re.sub(r"^\s*- \[[ xX]\] ", "", l), "[x]" in l.lower()[:8])
            elif re.match(r"^\s*[-*] ", l):
                self.vineta(re.sub(r"^\s*[-*] ", "", l), nivel=(len(l) - len(l.lstrip())) // 2)
            elif re.match(r"^\s*\d+\. ", l):
                self.vineta(re.sub(r"^\s*\d+\. ", "", l), numerada=True)
            elif l.strip() and not l.startswith("←"):
                self.p(l.strip())
            i += 1

    # ---------- salida ----------
    def guardar(self, ruta: Path, pdf: bool = True, con_comentarios: bool = False) -> list[Path]:
        ruta = Path(ruta); ruta.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(ruta)
        salida = [ruta]
        if pdf:
            salida.append(a_pdf(ruta, con_comentarios))
        return salida


def a_pdf(docx: Path, con_comentarios: bool = False) -> Path:
    """Exporta a PDF con Word (COM). Instancia propia (DispatchEx) y cierre garantizado."""
    import pythoncom, win32com.client
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False; word.DisplayAlerts = 0
    pdf = Path(docx).with_suffix(".pdf")
    try:
        doc = word.Documents.Open(str(Path(docx).resolve()), ReadOnly=True, AddToRecentFiles=False)
        doc.ExportAsFixedFormat(str(pdf.resolve()), 17, OptimizeFor=0, CreateBookmarks=1, DocStructureTags=True,
                                Item=7 if con_comentarios else 0)  # 7 = con marcas (comentarios en el margen)
        doc.Close(False)
    finally:
        word.Quit(); pythoncom.CoUninitialize()
    return pdf
