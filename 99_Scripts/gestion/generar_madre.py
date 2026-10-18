"""Genera el DOCUMENTO MADRE (00_DOCUMENTO_MADRE.docx + .pdf) y su versión para Obsidian (00_INICIO.md).

    python generar_madre.py

Lee: Cronograma_Tesis.xlsx (estado de tareas), 99_Scripts/datos/fuentes.sqlite (fuentes), Bitacora_Semanal.md
(alertas y pendientes), qué archivos existen en el repo. Lo ejecuta la tarea del domingo (tarea_domingo.py).
En el Word los enlaces van a GitHub (funcionan en cualquier PC); en 00_INICIO.md son enlaces de Obsidian.
"""
from __future__ import annotations

import re, sqlite3, subprocess, sys
from datetime import date
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
for s in (sys.stdout,):
    try: s.reconfigure(encoding="utf-8")
    except Exception: pass
from estilo_docx import Doc  # noqa: E402
import diagramas as dg  # noqa: E402
from generar_cronograma import ENT, SPRINTS, XLSX  # noqa: E402

BOVEDA = AQUI.parents[1]
GH = "https://github.com/l333x/Tesis_Torvesa"
IMG = BOVEDA / "11_Planificacion" / "Informes" / "_figuras"
HOY = date.today()


def rastreados() -> set[str]:
    r = subprocess.run(["git", "ls-files"], cwd=BOVEDA, capture_output=True, text=True, encoding="utf-8")
    return set(x for x in r.stdout.split("\n") if x)


TRACK = rastreados()
_MD = (BOVEDA / "11_Planificacion" / "CRONOGRAMA.md").read_text(encoding="utf-8")
NUM = {k: int(v) for k, v in re.findall(r"\*\*([A-Z]+\d*-\d+)\*\* .*· #(\d+)\s*$", _MD, re.M)}


def enlace(ruta: str, texto: str | None = None) -> str:
    """Enlace a GitHub si el archivo/carpeta está en el repo; si no, lo dice."""
    ruta = ruta.rstrip("/")
    texto = texto or ruta
    if ruta in TRACK:
        return f"[{texto}]({GH}/blob/main/{ruta})"
    if any(t.startswith(ruta + "/") for t in TRACK):
        return f"[{texto}]({GH}/tree/main/{ruta})"
    if (BOVEDA / ruta).exists():
        return f"`{texto}` (solo en tu PC)"
    return f"`{texto}` (aún no existe)"


def estados() -> dict:
    """Lee el Excel: tareas por sprint con su estado."""
    from openpyxl import load_workbook
    out = {}
    if not XLSX.exists():
        return out
    wb = load_workbook(XLSX, data_only=False)
    for clave, *_ in SPRINTS:
        if clave not in wb.sheetnames:
            continue
        filas = []
        for row in wb[clave].iter_rows(min_row=6, values_only=True):
            if row[3]:
                f = row[1].date() if hasattr(row[1], "date") else row[1]
                filas.append({"id": row[0], "fecha": f, "tarea": row[3], "resp": row[4], "donde": row[5], "estado": row[10] or "Pendiente"})
        out[clave] = filas
    return out


def sprint_actual():
    for clave, titulo, objetivo, tareas in SPRINTS:
        fechas = [t["fecha"] for t in tareas]
        ini, fin = min(fechas), max(fechas)
        lunes = ini - __import__("datetime").timedelta(days=ini.weekday())
        if lunes <= HOY <= fin:
            return clave, titulo, objetivo
        if HOY < lunes:
            return clave, titulo + " (empieza el lunes)", objetivo
    return SPRINTS[-1][0], SPRINTS[-1][1], SPRINTS[-1][2]


def fuentes() -> dict:
    db = BOVEDA / "99_Scripts" / "datos" / "fuentes.sqlite"
    if not db.exists():
        return {}
    c = sqlite3.connect(db)
    try:
        return {"total": c.execute("SELECT count(*) FROM fuentes").fetchone()[0],
                "por_estado": dict(c.execute("SELECT estado, count(*) FROM fuentes GROUP BY estado").fetchall()),
                "por_cat": c.execute("SELECT categoria, count(*), sum(estado='verificada') FROM fuentes GROUP BY categoria").fetchall()}
    except sqlite3.Error:
        return {}


def bitacora() -> tuple[str, dict]:
    t = (BOVEDA / "11_Planificacion" / "Bitacora_Semanal.md").read_text(encoding="utf-8")
    secs = [m for m in re.finditer(r"^## (\d{4}-\d{2}-\d{2}.*)$", t, re.M)]
    if not secs:
        return "", {}
    m = secs[-1]; cuerpo = t[m.end():]
    sub, act = {}, None
    for l in cuerpo.split("\n"):
        if l.startswith("## "):
            break
        h = re.match(r"^### (.+)", l)
        if h:
            act = h.group(1).strip(); sub[act] = []
        elif act and re.match(r"^\s*[-*] ", l):
            sub[act].append(re.sub(r"^\s*[-*] ", "", l).strip())
    return m.group(1), sub


CARPETAS = [
    ("00_Inbox", "Todo lo nuevo sin clasificar", "Ambos"), ("01_Borrador_Aprobado", "Diseño aprobado (no se edita)", "Nadie"),
    ("02_Fuentes", "Una ficha por fuente", "Víctor"), ("03_Matriz_Objetivos", "Afirmaciones ↔ fuentes ↔ estado", "Víctor"),
    ("04_Instrumentos", "Encuestas, entrevista, consentimiento", "Víctor"), ("05_Datos_Campo", "Datos reales anonimizados", "Ambos"),
    ("06_Capitulos", "Capítulos y modelo de negocio", "Víctor"), ("07_Flujos_Diseno", "Procesos y diseño del sistema", "Diego"),
    ("08_Reuniones_Tutora", "Actas con la tutora", "Víctor"), ("09_Registro_Contribuciones", "Quién hizo qué", "Ambos"),
    ("10_Antecedentes_Propios", "Trabajos previos de ustedes", "Ambos"), ("11_Planificacion", "Cronograma, informes, bitácora", "Víctor"),
    ("99_Scripts", "Programas (búsqueda, RAG, informes)", "Diego"),
]
CAPITULOS = [  # (título, qué contendrá, ids de entregable, categorías de fuentes, datos, flujos, patrón de archivo)
    ("Capítulo 1 · El problema", "Contexto de Torvesa, el problema (tiempo de respuesta, historial disperso), justificación y objetivos.",
     ["S2-01"], ["Problema", "Contexto"], "Línea base real (si hay campo) para describir el problema", "—", "Cap1"),
    ("Capítulo 2 · Marco teórico", "CRM y CRM móvil, tiempo de respuesta, gestión del conocimiento, IA con búsqueda en documentos (RAG), normativa.",
     ["S2-02"], ["Evidencia_Datos", "Tecnologia_IA", "Normativa"], "—", "—", "Cap2"),
    ("Capítulo 3 · Metodología", "Enfoque mixto con predominio cuantitativo, DSR (investigación en ciencia del diseño), población (dos censos), instrumentos, análisis y ética.",
     ["S2-03"], ["Metodologia"], "Instrumentos de 04_Instrumentos", "—", "Cap3"),
    ("Capítulo 4 · Diagnóstico", "Cómo funciona hoy la atención comercial de Torvesa, con datos reales.",
     ["S7-04"], ["Problema"], "05_Datos_Campo/LineaBase_2026-10", "Proceso actual (AS-IS)", "Cap4"),
    ("Capítulo 5 · Diseño del sistema y viabilidad", "Requisitos, procesos con el sistema (TO-BE), arquitectura, herramientas, costos, VAN y TIR, riesgos, ley de datos.",
     ["S4-04", "S8-06"], ["Tecnologia_IA", "Casos_Estudio", "Costos_Viabilidad", "Normativa"], "—", "TO-BE, bóveda del sistema, modelo de datos", "Cap5"),
    ("Capítulo 6 · Resultados de la validación", "Comparación del antes y el después de la validación del diseño en operación experimental controlada (metas M1–M6).",
     ["S9-02"], ["Evidencia_Datos"], "Línea base + medición «después»", "Casos de prueba", "Cap6"),
]
DECISIONES = [
    ("D1", "Lo que el borrador llama «piloto» se llama «validación del diseño en operación experimental controlada»; nunca «implementación»", "Tomada"),
    ("D2", "Marco DSR + Design Thinking para diseñar + Scrum para organizar; enfoque «mixto con predominio cuantitativo»", "Tomada (tutora sin objeciones, 12-oct; falta el acta)"),
    ("D3", "Usar los objetivos de aprendizaje oficiales de la carrera (OVA 4.1–6.3)", "Tomada"),
    ("D4", "Dos censos: los 30 colaboradores y los 33 clientes", "Tomada"),
    ("D5", "Entrega en enero (escenario B), salvo que la tutora confirme el 15-dic como fecha obligatoria (escenario A)", "Falta la respuesta de la tutora"),
    ("D6", "Proponer quitar OpenTimestamps (sello en Bitcoin) y usar huella SHA-256 + registro que solo se puede ampliar", "Tomada (tutora sin objeciones, 12-oct; falta el acta)"),
    ("D7", "VAN y TIR del sistema para Torvesa; el negocio SaaS TEAMSLASH va como anexo", "Tomada"),
]


DIAS = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def dia_es(f) -> str:
    return f"{DIAS[f.weekday()]} {f.day}-{MESES[f.month - 1]}" if hasattr(f, "weekday") else str(f)


def generar() -> None:
    est = estados(); clave_act, titulo_act, obj_act = sprint_actual(); fu = fuentes(); sem, bit = bitacora()
    tareas_act = est.get(clave_act, [])
    hechas = sum(1 for f in est.values() for t in f if t["estado"] == "Hecho"); total = sum(len(f) for f in est.values())
    d = Doc("Documento madre", "La tesis en un solo lugar · se actualiza cada domingo", f"Actualizado el {HOY.isoformat()}")

    # 1. Proyecto
    d.h1("1. Qué es el proyecto")
    d.p("**Tesis:** «Diseño de un sistema de automatización para la gestión comercial en la Constructora Torvesa» · Proyecto Técnico · "
        "Negocios Digitales, Universidad Politécnica Salesiana (UPS), Quito · Período 69 (sep 2026 – feb 2027) · "
        "**Autores:** Víctor Segura y Diego Bermeo · **Tutora:** Ing. Tania Chicaiza.")
    d.p("**El problema, en simple:** cuando un cliente pide una cotización, el vendedor la arma casi desde cero, buscando en archivos dispersos y en su memoria. "
        "Como pasa el día entre la oficina y la obra, las solicitudes esperan horas o días, y lo que sabe de cada cliente se queda en su celular.")
    d.p("**La solución que se diseña:** sobre el CRM TEAMSLASH (el sistema de clientes que ustedes ya hicieron) se agregan una **bóveda de conocimiento** "
        "(fichas de clientes, plantillas y cotizaciones anteriores), un **asistente de IA** que propone borradores de cotización y un **registro de eventos** "
        "que no se puede alterar. Se prueba en una **validación del diseño en operación experimental controlada** con el equipo de Torvesa.")
    d.imagen(dg.arquitectura(IMG / "arquitectura.png"), 16, "Figura 1. El sistema que se diseña, en una imagen")
    d.p("**Qué SÍ incluye:** diagnóstico con datos reales, diseño, prueba de concepto, validación controlada, costos, VAN y TIR, cumplimiento de la ley de datos (LOPDP). "
        "**Qué NO incluye:** instalar el sistema de forma definitiva, negociación, facturación, obra, contratación pública, ni los datos simulados del proyecto integrador.")

    # 2. Hoy
    d.salto(); d.h1("2. Dónde estamos hoy")
    d.tabla([["Semana actual", "Avance general", "Fuentes en la base", "Última bitácora"],
             [titulo_act, f"{hechas} de {total} tareas hechas ({(hechas / total * 100 if total else 0):.0f} %)",
              f"{fu.get('total', 0)} ({fu.get('por_estado', {}).get('verificada', 0)} verificadas)", sem or "—"]], anchos=[5, 4, 3.6, 4])
    hitos = [("12-oct", "Fase de fuentes"), ("14-oct", "Corte de fuentes"), ("15–16 oct", "Campo"), ("26–30 oct", "Revisión 1"),
             ("16–22 nov", "Revisión 2"), ("24-nov", "Validación"), ("18-dic", "Tesis al 90 %"), ("4–8 ene", "Entrega"), ("feb", "Defensa")]
    idx = {"S1": 0, "S2": 2, "S3": 3, "S4": 3, "S5": 4, "S6": 4, "S7": 5, "S8": 5, "S9": 6, "S10": 6, "Holgura": 7, "Entrega": 7, "Defensa": 8}.get(clave_act, 0)
    d.imagen(dg.linea_tiempo(hitos, IMG / "linea_tiempo.png", "Calendario (escenario B)", hoy=idx), 16.5)
    d.h2(f"Esta semana · {obj_act}")
    if tareas_act:
        d.tabla([["Día", "Tarea", "Quién", "Dónde", "Estado"]] +
                [[dia_es(t["fecha"]), t["tarea"], t["resp"], t["donde"], t["estado"]] for t in tareas_act],
                anchos=[2.1, 6.4, 1.6, 4.3, 2.2])
    for k, titulo in (("Riesgos y bloqueos", "Alertas"), ("Atención de Víctor", "Necesita tu decisión")):
        if bit.get(k):
            d.h2(titulo)
            for x in bit[k]:
                par = d.vineta(x)
                if k == "Atención de Víctor":
                    d.comentario(par, "Víctor: pendiente de tu decisión o acción.")

    # 3. Mapa
    d.salto(); d.h1("3. Mapa de la bóveda")
    d.p("La bóveda es la carpeta compartida de la tesis (`F:\\Obsidian\\Tesis_Torvesa`), guardada también en GitHub. Colores: azul = Víctor, cobre = Diego, azul oscuro = ambos, gris = nadie.")
    d.imagen(dg.mapa_carpetas(CARPETAS, IMG / "mapa.png"), 16.5)
    d.tabla([["Carpeta (clic para abrir)", "Para qué", "Responsable"]] + [[enlace(c), p, r] for c, p, r in CARPETAS], anchos=[6, 7.4, 3.2])
    d.p("**Documentos clave:** " + " · ".join([enlace("Glosario.pdf", "Glosario"), enlace("GUIA_DIEGO.pdf", "Guía de Diego"),
                                              enlace("Conceptos_Negocio.pdf", "Conceptos de negocio"), enlace("11_Planificacion/Cronograma_Tesis.xlsx", "Cronograma (Excel)"),
                                              enlace("99_Scripts/busqueda/Guia_Busqueda.pdf", "Guía de búsqueda"), enlace("11_Planificacion/Auditoria_Boveda2.pdf", "Auditoría")]))

    # 4. Capítulos
    d.salto(); d.h1("4. La tesis capítulo por capítulo")
    d.p("La estructura es provisional hasta que la carrera entregue el formato oficial.")
    for titulo, que, ids, cats, datos, flujos, patron in CAPITULOS:
        d.h2(titulo)
        d.p(f"**Qué contendrá:** {que}")
        arch = sorted(p.relative_to(BOVEDA).as_posix() for p in (BOVEDA / "06_Capitulos").glob(f"*{patron}*"))
        sem_ent = ", ".join(f"{i.split('-')[0]}" for i in ids)
        n_f = sum(n for c, n, _ in fu.get("por_cat", []) if c in cats); n_v = sum((v or 0) for c, _, v in fu.get("por_cat", []) if c in cats)
        d.tabla([["Aspecto", "Estado"],
                 ["Archivo del capítulo", " · ".join(enlace(a) for a in arch) if arch else f"aún no existe — se crea en el sprint {sem_ent} ({', '.join(f'issue #{ENT[i][0] or NUM.get(i)}' for i in ids if ENT[i][0] or NUM.get(i))})"],
                 ["Fuentes que lo respaldan", f"{n_f} fichas en {', '.join(cats)} ({n_v} verificadas) · " + enlace("02_Fuentes", "ver fichas")],
                 ["Datos de campo que usa", datos], ["Flujos de diseño que usa", flujos]], anchos=[4.2, 12.4])

    # 5. Fuentes
    d.h1("5. Fuentes")
    if fu:
        d.tabla([["Categoría", "Fichas", "Verificadas"]] + [[c, str(n), str(v or 0)] for c, n, v in fu["por_cat"]], anchos=[6, 3, 3])
    else:
        d.p("Todavía no hay fuentes en la base: la fase de fuentes empieza el lunes 12-oct.")
    d.p("**Matriz afirmación–fuente–estado** (la tabla que dice qué afirmación de la tesis tiene respaldo): " +
        enlace("03_Matriz_Objetivos/2026-10-14_Matriz_Afirmaciones_v0.1.xlsx", "matriz") + " — se crea en el corte del miércoles 14-oct. "
        "**Base completa:** " + enlace("99_Scripts/datos/fuentes.csv", "fuentes.csv") + ". Solo se cita lo «verificado» (abierto, con página).")

    # 6. Campo
    d.h1("6. Datos de campo")
    d.tabla([["Qué se levanta", "Cuándo", "Dónde queda"],
             ["Encuesta a colaboradores (I1) y entrevistas (I3)", "15–16 oct, si Torvesa autoriza", enlace("05_Datos_Campo/LineaBase_2026-10")],
             ["Encuesta a clientes por teléfono (I2)", "16-oct, si Torvesa autoriza", enlace("05_Datos_Campo/LineaBase_2026-10")],
             ["Tiempo de cada cotización, con cronómetro (P2-A)", "15–16 oct", enlace("05_Datos_Campo/LineaBase_2026-10")],
             ["Fechas y horas de solicitudes pasadas (P2-B)", "16-oct", enlace("05_Datos_Campo/LineaBase_2026-10")],
             ["Medición «después» (mismos instrumentos)", "7–9 dic", "05_Datos_Campo/ (carpeta Post, se crea en diciembre)"],
             ["Datos con nombres o teléfonos", "Siempre", "05_Datos_Campo/_crudo/ — solo en tu PC, nunca en GitHub"]], anchos=[6.6, 4.2, 5.8])
    d.aviso("Los datos del «diagnóstico N = 30» del proyecto integrador eran **simulados**: no se usan. " + enlace("05_Datos_Campo/Diagnostico_2026_SIMULADO_no_usar/LEEME_Procedencia.md", "Detalle"), "alerta")

    # 7. Diseño
    d.h1("7. Diseño del sistema")
    fl = sorted(p.relative_to(BOVEDA).as_posix() for p in (BOVEDA / "07_Flujos_Diseno").glob("*") if p.name != "README.md")
    d.p("Archivos de diseño: " + (" · ".join(enlace(a) for a in fl) if fl else "aún no hay — el proceso actual (AS-IS) se dibuja en el sprint 2 y el proceso con el sistema (TO-BE) en el sprint 4."))

    # 8. Finanzas
    d.h1("8. Evaluación financiera y modelo de negocio")
    d.tabla([["Pieza", "Qué es", "Cuándo"],
             ["Precios de herramientas con fuente", "Cuánto cuesta cada pieza (servidor, IA, licencias), con enlace y fecha", "Sprint 1 · Diego"],
             ["Lienzo de modelo de negocio y propuesta de valor", "A quién sirve, qué valor da, cómo se sostiene", "Sprint 2 · Víctor"],
             ["CAPEX y OPEX", "CAPEX = inversión inicial · OPEX = gasto de cada mes", "Sprint 4 · Víctor"],
             ["Beneficios con supuestos explícitos", "Horas ahorradas al cotizar, clientes que no se pierden", "Sprint 5 · Víctor"],
             ["VAN, TIR y período de recuperación", "Si conviene invertir, en 3 escenarios (pesimista, base, optimista), con sensibilidad", "Sprint 5 · Víctor"],
             ["Pitch de 1 página para inversionista", "La idea y los números en una hoja", "Sprint 5 · Víctor"]], anchos=[5.4, 7.4, 3.8])
    mf = sorted(p.relative_to(BOVEDA).as_posix() for p in (BOVEDA / "06_Capitulos" / "Modelo_Negocio").glob("*") if p.name != "README.md")
    d.p("Archivos: " + (" · ".join(enlace(a) for a in mf) if mf else "aún no existen — la carpeta es " + enlace("06_Capitulos/Modelo_Negocio") + ". El modelo irá en Excel con fórmulas vivas, formato condicional y notas al lado de cada supuesto."))

    # 9. Decisiones
    d.h1("9. Decisiones y pendientes")
    d.tabla([["Código", "Decisión", "Estado"]] + [list(x) for x in DECISIONES], anchos=[1.6, 11.4, 3.6])
    d.p("**Pendientes con Torvesa (pedido hasta el 14-oct):** confirmar a qué se dedica, nota de respaldo, catálogo de servicios y permisos de campo. "
        "**Con la tutora:** datos simulados, fecha 15-dic, D2, D6 y duración de la validación — " + enlace("08_Reuniones_Tutora/2026-10-12_Puntos_Tutora.pdf", "hoja de puntos") + ".")

    # 10. Glosario
    d.salto(); d.h1("10. Glosario")
    glo = (BOVEDA / "Glosario.md").read_text(encoding="utf-8")
    glo = re.sub(r"^---\n.*?\n---\n", "", glo, count=1, flags=re.S); glo = re.sub(r"^# .*\n", "", glo, count=1, flags=re.M)
    glo = re.sub(r"^## ", "### ", glo, flags=re.M)
    d.markdown(glo)
    d.guardar(BOVEDA / "00_DOCUMENTO_MADRE.docx", pdf=True)

    # ---- versión Obsidian ----
    L = ["---", "titulo: Inicio — documento madre (versión Obsidian)", f"actualizado: {HOY.isoformat()}", "tipo: portal", "tags: [inicio, portal]", "---", "",
         "# Inicio", "", f"> Versión para leer en Obsidian. La versión completa, con gráficos, es [[00_DOCUMENTO_MADRE.pdf]] (se regenera cada domingo).", "",
         "## Hoy", f"- **Semana:** {titulo_act}", f"- **Avance:** {hechas}/{total} tareas · **Fuentes:** {fu.get('total', 0)} ({fu.get('por_estado', {}).get('verificada', 0)} verificadas)",
         "- **Cronograma:** [[Cronograma_Tesis.xlsx]] · [[CRONOGRAMA]] (versión IA)", ""]
    if tareas_act:
        L += ["### Esta semana", "| Día | Tarea | Quién | Estado |", "|---|---|---|---|"]
        L += [f"| {dia_es(t['fecha'])} | {t['tarea']} | {t['resp']} | {t['estado']} |" for t in tareas_act]
    if bit.get("Atención de Víctor"):
        L += ["", "### Necesita tu decisión"] + [f"- [ ] {x}" for x in bit["Atención de Víctor"]]
    L += ["", "## Carpetas"] + [f"- [[{c}/README|{c}]] — {p} ({r})" for c, p, r in CARPETAS]
    L += ["", "## Documentos clave", "- [[Glosario]] · [[GUIA_DIEGO]] · [[Conceptos_Negocio.pdf]] · [[Guia_Busqueda.pdf]] · [[Auditoria_Boveda2.pdf]]",
          "- [[Bitacora_Semanal]] · [[Registro_Contribuciones]] · [[2026-10-11_F0_Validacion_Tesis|Fase 0]]", "", "## Capítulos"]
    for titulo, que, ids, *_r, patron in CAPITULOS:
        arch = sorted(p.stem for p in (BOVEDA / "06_Capitulos").glob(f"*{patron}*"))
        L.append(f"- **{titulo}** — {que} " + (" ".join(f"[[{a}]]" for a in arch) if arch else f"*(aún no existe — sprint {ids[0].split('-')[0]})*"))
    L += ["", "## Decisiones"] + [f"- **{c}** {t} — *{e}*" for c, t, e in DECISIONES]
    (BOVEDA / "00_INICIO.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    print("ok 00_DOCUMENTO_MADRE.docx/.pdf + 00_INICIO.md")


if __name__ == "__main__":
    generar()
